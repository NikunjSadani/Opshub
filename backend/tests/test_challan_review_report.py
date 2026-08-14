"""The downloadable Excel review report for a NEEDS_REVIEW batch: two sheets, the
uploaded rows echoed with a status flag, and one line per contradiction."""
from __future__ import annotations

import io

from openpyxl import load_workbook

from app.modules.challan import review_report
from app.modules.challan.schema import CHALLAN_COLUMNS, COLUMN_HEADERS, Contradiction, RawRow


def _row(n: int, gstin: str, name: str) -> RawRow:
    cells = {k: "" for k in CHALLAN_COLUMNS}
    cells["consignee_gstin"] = gstin
    cells["consignee_name"] = name
    cells["challan_group"] = f"G{n}"
    return RawRow(row_number=n, cells=cells)


def test_review_report_has_two_sheets_and_flags_contradictions() -> None:
    rows = [_row(2, "24AAACB2894G1ZT", "Deoleo India Ltd"),
            _row(3, "27AAPFU0939F1ZV", "Clean Co")]
    contradictions = [Contradiction(
        gstin="24AAACB2894G1ZT", consignee_name="Deoleo India Ltd",
        field="name", stored="Deoleo MH", uploaded="Deoleo India Ltd")]

    data = review_report.build_review_xlsx(rows, contradictions)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Uploaded rows", "Contradictions to review"]

    uploaded = wb["Uploaded rows"]
    header = [c.value for c in next(uploaded.iter_rows(max_row=1))]
    # Row + every uploaded column (friendly header) + a status column.
    assert header[0] == "Row"
    assert header[1:1 + len(CHALLAN_COLUMNS)] == [COLUMN_HEADERS[k] for k in CHALLAN_COLUMNS]
    assert header[-1] == "Review status"
    statuses = [r[-1].value for r in uploaded.iter_rows(min_row=2)]
    assert any("Contradiction" in (s or "") for s in statuses)
    assert "OK" in statuses  # the clean row

    review = wb["Contradictions to review"]
    body = [[c.value for c in r] for r in review.iter_rows(min_row=2)]
    assert body and body[0][0] == "24AAACB2894G1ZT"
    assert body[0][3] == "Deoleo MH" and body[0][4] == "Deoleo India Ltd"


def test_review_report_neutralizes_formula_injection() -> None:
    # An untrusted upload cell / consignee value beginning with '=' must NOT become a
    # live Excel formula (openpyxl stores a leading '=' as data_type 'f').
    rows = [_row(2, "24AAACB2894G1ZT", "=cmd|'/c calc'!A1")]
    contradictions = [Contradiction(
        gstin="24AAACB2894G1ZT", consignee_name="=cmd|'/c calc'!A1",
        field="name", stored="Deoleo MH", uploaded="=HYPERLINK(0)")]
    data = review_report.build_review_xlsx(rows, contradictions)
    wb = load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.data_type != "f", (ws.title, cell.coordinate, cell.value)
    # the neutralized name kept its text (prefixed with a quote), not evaluated
    review = wb["Contradictions to review"]
    assert any(str(c.value).startswith("'=") for r in review.iter_rows() for c in r)
