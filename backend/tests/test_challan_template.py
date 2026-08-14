"""The self-documenting upload template: structure, the download route, and a
ROUND-TRIP proof that the template's own example rows parse + validate cleanly
(structurally), so the examples can never silently rot."""
from __future__ import annotations

import io
from collections.abc import Iterator

import pytest
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool
from starlette.testclient import TestClient

from app.db import Base, get_db
from app.modules.challan import parsing, template
from app.modules.challan.routes import router
from app.modules.challan.schema import CHALLAN_COLUMNS, COLUMN_HEADERS
from app.platform.auth import current_user
from app.platform.models import Role, User, UserModuleAccess

MIS = User(firebase_uid="mis", email="m@x.com", name="M", role=Role.MIS, active=True,
           module_access=[UserModuleAccess(module_key="document_automation")])
OUTSIDER = User(firebase_uid="out", email="o@x.com", name="O", role=Role.OPERATIONS, active=True)


@pytest.fixture
def client() -> Iterator[TestClient]:
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False},
                           poolclass=StaticPool, future=True)
    Base.metadata.create_all(engine)
    TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False, future=True)

    def _db() -> Iterator[Session]:
        db = TestSession()
        try:
            yield db
        finally:
            db.close()

    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.dependency_overrides[get_db] = _db
    app.dependency_overrides[current_user] = lambda: MIS
    yield TestClient(app)
    engine.dispose()


# --------------------------------------------------------------------- tests

def test_template_structure() -> None:
    wb = load_workbook(io.BytesIO(template.build_template_xlsx()))
    # The data sheet MUST be first (the parser reads worksheets[0]).
    assert wb.sheetnames[0] == "Challans"
    assert "Instructions" in wb.sheetnames
    ws = wb.worksheets[0]
    header = [c.value for c in ws[1]]
    # Row 1 is written with the FRIENDLY display headers, in canonical column order.
    assert header == [COLUMN_HEADERS[key] for key in CHALLAN_COLUMNS]


def test_template_examples_roundtrip_validate() -> None:
    # The template's OWN example rows must parse + pass the no-database checks.
    rows, structural = parsing.parse_workbook(template.build_template_xlsx())
    assert structural == []  # friendly headers all recognised, no missing columns
    assert parsing.structural_row_errors(rows) == []
    # ...and they group into the two example challans by the Challan Group column.
    groups: dict[str, list[object]] = {}
    for row in rows:
        groups.setdefault(row.cells["challan_group"], []).append(row)
    assert set(groups) == {"C1", "C2"}
    assert len(groups["C1"]) == 1   # single priced line
    assert len(groups["C2"]) == 2   # two value-free lines


def test_template_download_route(client: TestClient) -> None:
    r = client.get("/api/v1/challan/template.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "challan-upload-template.xlsx" in r.headers["content-disposition"]
    # the bytes are a real workbook
    assert load_workbook(io.BytesIO(r.content)).sheetnames[0] == "Challans"


def test_template_download_requires_module(client: TestClient) -> None:
    client.app.dependency_overrides[current_user] = lambda: OUTSIDER
    assert client.get("/api/v1/challan/template.xlsx").status_code == 403
