"""Project-Product tag service — curates a project's PO product picker + pricing template.

A tag is a link row (`project_product`) between a SHARED Product master and a Project,
never a re-parenting: one product can be tagged to many projects and vice-versa. The
`uq_project_product` unique constraint is the DB backstop, so the tag/untag operations
are IDEMPOTENT — a re-tag returns the existing row (WITHOUT overwriting its fields), an
un-tag of an absent pair is a no-op. Every write is audited inside the caller's
transaction (WoW: everything traceable).

Each tag also carries a per-project PRICING TEMPLATE (mirrors the PO line, minus qty) a
PO pre-fills from. `sell_price_paise`/`freight_paise` are the ACTUAL (margin-basis)
figures — ADMIN-ONLY; the ROUTE masks them on read and drops them from a non-admin's
write, so this service applies only what it is handed.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session, joinedload

from app.modules.challan.parsing import parse_paise, parse_qty
from app.modules.projects.models import Project
from app.modules.sales_orders.models import Product, ProjectProduct
from app.platform import audit

# Money ceiling (mirrors the route's ProjectProductWrite bound + the billing money guard):
# a per-cell paise value beyond this degrades to a per-row error, never an int8-overflow 500.
_MAX_MONEY_PAISE = 10**15
# Column caps mirrored from the single-write path (ProjectProductWrite / the DB columns).
_MAX_DESCRIPTION_LEN = 500
_MAX_UOM_LEN = 20
# The two ADMIN-ONLY actual template fields — dropped from a non-admin's bulk write (below),
# exactly as the single-PATCH route drops them via `_drop_actuals_if_masked`.
_ACTUAL_FIELD_KEYS: tuple[str, ...] = ("sell_price_paise", "freight_paise")

# The pricing-template columns a tag/update may carry (mirror the PO line, no ordered_qty).
# The route validates bounds (money ge 0, tax 0..100) and drops the two ADMIN-ONLY actuals
# (`sell_price_paise`/`freight_paise`) for a non-admin BEFORE calling in — the service just
# applies whatever survives.
PRICING_FIELDS: tuple[str, ...] = (
    "description",
    "uom",
    "cost_price_paise",
    "original_cost_price_paise",
    "client_sell_price_paise",
    "vendor_sell_price_paise",
    "sell_price_paise",
    "client_freight_paise",
    "vendor_freight_paise",
    "freight_paise",
    "packaging_paise",
    "handling_paise",
    "other_paise",
    "tax_rate",
)


class ProjectProductError(Exception):
    """Base error for a project-product tag operation."""


class ProjectNotFound(ProjectProductError):
    """The referenced project is missing (route -> 404)."""


class ProductNotFound(ProjectProductError):
    """The referenced product is missing (route -> 404)."""


class TagNotFound(ProjectProductError):
    """The (project, product) tag being edited does not exist (route -> 404)."""


def _require_project(db: Session, project_id: int) -> Project:
    project = db.get(Project, project_id)
    if project is None:
        raise ProjectNotFound(f"project {project_id} not found")
    return project


def _apply_fields(tag: ProjectProduct, fields: dict[str, Any]) -> None:
    """Set only the recognized pricing keys present in `fields` on `tag`. Unknown keys
    are ignored (the route's pydantic model is the field allow-list)."""
    for key in PRICING_FIELDS:
        if key in fields:
            setattr(tag, key, fields[key])


def list_tagged(db: Session, *, project_id: int) -> list[ProjectProduct]:
    """Return every tag row for `project_id` with its joined product (including inactive
    products — the FE marks them), ordered by product name. Raises `ProjectNotFound` if
    the project doesn't exist. Each row carries the per-project pricing template."""
    _require_project(db, project_id)
    stmt = (
        select(ProjectProduct)
        .join(Product, ProjectProduct.product_id == Product.id)
        .options(joinedload(ProjectProduct.product))
        .where(ProjectProduct.project_id == project_id)
        .order_by(Product.name)
    )
    return list(db.execute(stmt).scalars())


def tag_product(
    db: Session,
    *,
    project_id: int,
    product_id: int,
    fields: dict[str, Any] | None = None,
    actor_uid: str | None = None,
) -> tuple[ProjectProduct, bool]:
    """Tag `product_id` to `project_id`, optionally seeding the pricing template from
    `fields`. IDEMPOTENT: if the pair is already tagged, return the EXISTING tag with
    created=False and DO NOT overwrite its fields (editing is via `update_template`);
    otherwise create the tag with `fields` and return created=True. Validates both the
    project and the product exist (404 each). Audited `sales_orders.product_tagged` on a
    NEW tag only.

    Returns `(tag, created)`.
    """
    _require_project(db, project_id)
    product = db.get(Product, product_id)
    if product is None:
        raise ProductNotFound(f"product {product_id} not found")

    existing = db.execute(
        select(ProjectProduct).where(
            ProjectProduct.project_id == project_id,
            ProjectProduct.product_id == product_id,
        )
    ).scalar_one_or_none()
    if existing is not None:
        # Idempotent re-tag: never clobber an already-templated row.
        return existing, False

    tag = ProjectProduct(
        project_id=project_id, product_id=product_id, created_by=actor_uid
    )
    seed = dict(fields or {})
    # Default the template's tax rate from the product's GST rate on a NEW tag only, so the
    # rate flows product -> project template -> PO without re-typing. Only when NO usable
    # tax_rate was supplied (key absent or None) — an explicit value (incl. Decimal('0') for
    # a 0% nil-rated line) is always honoured, never overridden.
    if seed.get("tax_rate") is None and product.gst_rate is not None:
        seed["tax_rate"] = product.gst_rate
    _apply_fields(tag, seed)
    try:
        # `with` releases the savepoint on success; a concurrent tag of the same pair
        # trips the `uq_project_product` unique constraint at flush and lands here —
        # idempotent, so the pre-existing tag is treated as success.
        with db.begin_nested():
            db.add(tag)
            db.flush()
    except IntegrityError:
        # Lost the race: the pair now exists — return it without overwriting its fields.
        winner = db.execute(
            select(ProjectProduct).where(
                ProjectProduct.project_id == project_id,
                ProjectProduct.product_id == product_id,
            )
        ).scalar_one()
        return winner, False

    audit.log(
        db,
        action="sales_orders.product_tagged",
        actor_uid=actor_uid,
        entity="project_product",
        entity_id=str(tag.id),
        detail={"project_id": project_id, "product_id": product_id},
    )
    return tag, True


def update_template(
    db: Session,
    *,
    project_id: int,
    product_id: int,
    fields: dict[str, Any],
    actor_uid: str | None = None,
) -> ProjectProduct:
    """Partial-update the pricing template of an EXISTING (project, product) tag: apply
    only the keys present in `fields` (the route passes an `exclude_unset` dump, already
    stripped of the ADMIN-ONLY actuals for a non-admin), leaving every other column
    untouched. Raises `TagNotFound` (route -> 404) if the pair isn't tagged.

    Audited `sales_orders.product_template_updated` with the SORTED changed field NAMES
    only — never the actual money values — so the trail can't leak a masked actual.
    """
    tag = db.execute(
        select(ProjectProduct).where(
            ProjectProduct.project_id == project_id,
            ProjectProduct.product_id == product_id,
        )
    ).scalar_one_or_none()
    if tag is None:
        raise TagNotFound(
            f"product {product_id} is not tagged to project {project_id}"
        )

    changed = [key for key in PRICING_FIELDS if key in fields]
    _apply_fields(tag, fields)
    db.flush()

    audit.log(
        db,
        action="sales_orders.product_template_updated",
        actor_uid=actor_uid,
        entity="project_product",
        entity_id=str(tag.id),
        detail={
            "project_id": project_id,
            "product_id": product_id,
            "fields": sorted(changed),
        },
    )
    return tag


def untag_product(
    db: Session,
    *,
    project_id: int,
    product_id: int,
    actor_uid: str | None = None,
) -> bool:
    """Remove the (project, product) tag if present. IDEMPOTENT: returns True when a
    row was removed, False when there was nothing to remove. Audited
    `sales_orders.product_untagged` only when a row was removed."""
    tag = db.execute(
        select(ProjectProduct).where(
            ProjectProduct.project_id == project_id,
            ProjectProduct.product_id == product_id,
        )
    ).scalar_one_or_none()
    if tag is None:
        return False

    tag_id = tag.id  # capture before delete — a deleted instance's attrs may expire
    db.delete(tag)
    db.flush()
    audit.log(
        db,
        action="sales_orders.product_untagged",
        actor_uid=actor_uid,
        entity="project_product",
        entity_id=str(tag_id),
        detail={"project_id": project_id, "product_id": product_id},
    )
    return True


# ----------------------------------------------------------- bulk pricing (Excel)
#
# Bulk-set a PROJECT'S per-product pricing templates from an .xlsx: ONE sheet for ONE
# project (the project is chosen at UPLOAD time, like the PO upload chooses a client — a
# form field, not a sheet column). Each row identifies a product (by code, else by exact
# name) and carries the same pricing-template fields the single POST/PATCH accepts. Per
# row we TAG the product to the project (idempotent) then UPSERT its template.
#
# MASKING (mirrors the single-PATCH path exactly): the route computes `can_set_actuals`
# from the uploading user and passes it in — a non-admin's `actual_sell`/`actual_freight`
# (`sell_price_paise`/`freight_paise`) are DROPPED before persist, so a non-admin's sheet
# can never set the margin.


@dataclass
class PricingBulkResult:
    """Outcome of a bulk pricing upload: the product CODES successfully tagged+priced,
    and per-row errors (1-based sheet row + message)."""

    priced: list[str] = field(default_factory=list)
    errors: list[tuple[int, str]] = field(default_factory=list)


# Canonical bulk columns -> the friendly headers accepted for each (case-insensitive,
# whitespace/hyphen collapsed to '_'). `product_code` OR `product_name` must be present to
# identify the product; every pricing column is OPTIONAL (a blank cell leaves that template
# field untouched — this is an upsert, not a full replace).
_PRICING_ALIASES: dict[str, str] = {
    "product_code": "product_code", "code": "product_code", "sku": "product_code",
    "product_name": "product_name", "product": "product_name", "name": "product_name",
    "description": "description", "desc": "description",
    "uom": "uom", "unit": "uom",
    "cost_price": "cost_price", "cost": "cost_price",
    "original_cost_price": "original_cost_price", "original_cost": "original_cost_price",
    "client_sell": "client_sell", "client_sell_price": "client_sell",
    "vendor_sell": "vendor_sell", "vendor_sell_price": "vendor_sell",
    "client_freight": "client_freight",
    "vendor_freight": "vendor_freight",
    "packaging": "packaging",
    "handling": "handling",
    "other": "other",
    "tax_rate": "tax_rate", "tax": "tax_rate", "gst": "tax_rate", "gst_rate": "tax_rate",
    # ADMIN-ONLY actuals (the margin basis) — accepted here but dropped for a non-admin.
    "actual_sell": "actual_sell", "actual_sell_price": "actual_sell",
    "actual_freight": "actual_freight",
}

# Money column (rupees) -> the paise template field it sets. The two ADMIN-ONLY actuals map
# to `sell_price_paise`/`freight_paise` (dropped for a non-admin BEFORE persist).
_MONEY_COLUMN_TO_FIELD: dict[str, str] = {
    "cost_price": "cost_price_paise",
    "original_cost_price": "original_cost_price_paise",
    "client_sell": "client_sell_price_paise",
    "vendor_sell": "vendor_sell_price_paise",
    "client_freight": "client_freight_paise",
    "vendor_freight": "vendor_freight_paise",
    "packaging": "packaging_paise",
    "handling": "handling_paise",
    "other": "other_paise",
    "actual_sell": "sell_price_paise",
    "actual_freight": "freight_paise",
}

# The downloadable TEMPLATE header, in a fixed friendly order. The two ACTUAL columns are
# appended ONLY for an admin (a non-admin's template omits them — they could never set them).
# Money columns are RUPEES (converted to paise via `parse_paise`), so the example uses rupees.
_PRICING_TEMPLATE_COLUMNS: tuple[str, ...] = (
    "product_code", "description", "uom",
    "cost_price", "original_cost_price", "client_sell", "vendor_sell",
    "client_freight", "vendor_freight", "packaging", "handling", "other", "tax_rate",
)
_PRICING_TEMPLATE_ACTUAL_COLUMNS: tuple[str, ...] = ("actual_sell", "actual_freight")
# One example row aligned to _PRICING_TEMPLATE_COLUMNS (+ the two actuals when included).
# Money is in RUPEES (250.00 -> 25000 paise). product_code is illustrative — use a real SKU.
_PRICING_TEMPLATE_EXAMPLE: tuple[object, ...] = (
    "SKU-1001", "Sample line — replace with your product", "PCS",
    250.00, 260.00, 399.00, 380.00, 50.00, 40.00, 10.00, 20.00, 5.00, 18,
)
_PRICING_TEMPLATE_ACTUAL_EXAMPLE: tuple[object, ...] = (420.00, 55.00)


def build_pricing_template_xlsx(*, include_actuals: bool) -> bytes:
    """Return the bulk pricing-upload template as ``.xlsx`` bytes: the header row in the
    exact column order the parser accepts + one illustrative example row (money in rupees).
    The two ADMIN-ONLY actual columns are included only when ``include_actuals`` (admin)."""
    columns = list(_PRICING_TEMPLATE_COLUMNS)
    example = list(_PRICING_TEMPLATE_EXAMPLE)
    if include_actuals:
        columns += list(_PRICING_TEMPLATE_ACTUAL_COLUMNS)
        example += list(_PRICING_TEMPLATE_ACTUAL_EXAMPLE)
    wb = Workbook()
    ws = wb.active or wb.create_sheet()
    ws.title = "pricing"
    ws.append(columns)
    ws.append(example)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm_header(raw: str) -> str:
    return "_".join(raw.strip().lower().replace("-", " ").split())


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()


def _parse_pricing_workbook(
    data: bytes,
) -> tuple[list[tuple[int, dict[str, str]]], list[tuple[int, str]]]:
    """Read the first worksheet against the pricing template. Returns (rows, errors); a bad
    file or a missing product-key column yields ([], [errors]) so nothing is attempted."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises many types on a bad file
        return [], [(1, "could not read the file as an .xlsx workbook")]
    try:
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        col_index: dict[str, int] = {}
        for idx, raw in enumerate(header or ()):
            if raw is None:
                continue
            key = _PRICING_ALIASES.get(_norm_header(str(raw)))
            if key is not None and key not in col_index:
                col_index[key] = idx

        if "product_code" not in col_index and "product_name" not in col_index:
            return [], [(1, "a 'product_code' or 'product_name' column is required")]

        rows: list[tuple[int, dict[str, str]]] = []
        for sheet_row, values in enumerate(rows_iter, start=2):
            cells = {
                key: _cell_str(values[idx] if idx < len(values) else None)
                for key, idx in col_index.items()
            }
            if all(v == "" for v in cells.values()):
                continue  # skip a fully-blank row
            rows.append((sheet_row, cells))
        return rows, []
    finally:
        wb.close()


def _resolve_pricing_product(db: Session, code: str, name: str) -> Product | str:
    """Resolve the row's product by code (EXACT — the `code` unique is case-sensitive, so an
    exact match mirrors the Products-upload path and can never return >1 row) else by EXACT
    name (case-insensitive, ambiguity → error). Returns the Product, or an error STRING —
    never creates a product (a missing product must be added via the Products upload first)."""
    if code:
        product = db.execute(
            select(Product).where(Product.code == code)
        ).scalar_one_or_none()
        if product is None:
            return "product not found — add it via the Products upload first"
        return product
    if name:
        matches = list(db.execute(
            select(Product).where(func.lower(Product.name) == name.lower())
        ).scalars())
        if not matches:
            return "product not found — add it via the Products upload first"
        if len(matches) > 1:
            return "ambiguous product name; use a product code"
        return matches[0]
    return "a product_code or product_name is required"


def _build_pricing_row(
    db: Session, cells: dict[str, str]
) -> tuple[Product, dict[str, Any]] | str:
    """Turn one Excel row into ``(product, template_fields)`` or an error string. Money is
    parsed rupees->paise and bounded [0, ceiling]; tax is 0..100 and rejected over-scale
    (>2 dp) — same guards as the single-write path. A blank cell leaves that field unset."""
    product = _resolve_pricing_product(
        db, cells.get("product_code", "").strip(), cells.get("product_name", "").strip())
    if isinstance(product, str):
        return product

    fields: dict[str, Any] = {}
    for col, field_name in _MONEY_COLUMN_TO_FIELD.items():
        raw = cells.get(col, "").strip()
        if not raw:
            continue
        paise = parse_paise(raw)
        if paise is None or paise < 0:
            return f"{col} must be a number >= 0"
        if paise > _MAX_MONEY_PAISE:
            return f"{col} is too large"
        fields[field_name] = paise

    tax_raw = cells.get("tax_rate", "").strip()
    if tax_raw:
        parsed = parse_qty(tax_raw)
        if parsed is None or parsed < 0 or parsed > 100:
            return "tax_rate must be a number between 0 and 100"
        # Numeric(5,2): reject an over-scale tax rather than let Postgres silently round it.
        if parsed != parsed.quantize(Decimal("0.01")):
            return "tax_rate supports at most 2 decimal places"
        fields["tax_rate"] = parsed

    description = cells.get("description", "").strip()
    if description:
        fields["description"] = description[:_MAX_DESCRIPTION_LEN]
    uom = cells.get("uom", "").strip()
    if uom:
        fields["uom"] = uom[:_MAX_UOM_LEN]
    return product, fields


def bulk_set_pricing_from_excel(
    db: Session,
    data: bytes,
    *,
    project_id: int,
    can_set_actuals: bool,
    actor_uid: str | None = None,
) -> PricingBulkResult:
    """Parse an ``.xlsx`` and, for the given project, TAG each resolved product (idempotent)
    then UPSERT its pricing template. Row discipline (WoW: nothing silently dropped): a
    malformed / unresolved / ambiguous row lands in ``errors`` and is skipped; a resolved row
    is tagged+priced and its product CODE lands in ``priced``.

    MASKING: when ``can_set_actuals`` is False the two ADMIN-ONLY actuals
    (``sell_price_paise``/``freight_paise``) are DROPPED from every row before persist — a
    non-admin's sheet can never set the margin (mirrors the single-PATCH route).

    Each row's persist runs inside its OWN SAVEPOINT so one row's DB failure can't poison the
    batch. Raises ``ProjectNotFound`` (route -> 400) if the project doesn't exist. Caller
    commits."""
    _require_project(db, project_id)

    rows, parse_errors = _parse_pricing_workbook(data)
    result = PricingBulkResult(errors=list(parse_errors))
    if parse_errors:  # header/file problem — nothing to attempt
        return result

    for sheet_row, cells in rows:
        built = _build_pricing_row(db, cells)
        if isinstance(built, str):
            result.errors.append((sheet_row, built))
            continue
        product, fields = built
        # A non-admin can never persist an actual — drop both actual keys before persist.
        if not can_set_actuals:
            for key in _ACTUAL_FIELD_KEYS:
                fields.pop(key, None)
        try:
            # Per-row SAVEPOINT: a DB failure on this row rolls back only this row, leaving
            # the rest of the batch intact (the outer transaction is the caller's commit).
            with db.begin_nested():
                tag_product(
                    db, project_id=project_id, product_id=product.id,
                    fields=None, actor_uid=actor_uid)
                if fields:
                    update_template(
                        db, project_id=project_id, product_id=product.id,
                        fields=fields, actor_uid=actor_uid)
        except SQLAlchemyError:
            result.errors.append((sheet_row, "could not save this row"))
            continue
        # `code` is always populated (auto-minted PRD-###### when none supplied); guard for mypy.
        result.priced.append(
            product.code if product.code is not None else f"PRD-{product.id:06d}")

    return result
