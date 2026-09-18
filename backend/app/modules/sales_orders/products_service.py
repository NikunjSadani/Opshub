"""Product Master service — the normalized product identity PO lines reference.

Integrity discipline (mirrors the projects/numbering services):

  * A product's identity is `(lower(name), lower(coalesce(brand,'')),
    lower(coalesce(model_number,'')))`, enforced case-insensitively by the DB
    functional unique index `uq_product_identity`. The optional `code` is a
    second globally-unique key. Both are the real dedup backstop — a pre-check
    can't close the race, so the insert/update runs inside a SAVEPOINT and a
    collision at flush is caught and re-raised as `DuplicateProduct` (route
    -> 409).
  * Every mutation is audited inside the caller's transaction, so the trail
    commits atomically with the change.

⚠️ `db.begin_nested()` MUST be used as `with db.begin_nested():` — a bare call
leaks a SAVEPOINT per insert.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from dataclasses import field as dc_field
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.modules.challan.parsing import parse_qty
from app.modules.masterdata.models import HsnCode
from app.modules.sales_orders.models import Product, ProjectProduct
from app.platform import audit

# The auto-minted code namespace: `PRD-` + the zero-padded primary key (>= 6
# digits; ids past 999999 widen to 7+ digits). A user-supplied code in this
# shape is rejected (create + update) so no hand-entered code can ever collide
# with a system-generated one.
_RESERVED_CODE_RE = re.compile(r"^PRD-\d{6,}$", re.IGNORECASE)

# A valid statutory HSN/SAC code (mirrors the HSN master's own `HsnIn` guard): 2–12 digits.
# The product path only seeds/reconciles `md_hsn` for codes in this shape, so a product's
# free-form `hsn` (e.g. "8471 3090" or "ABC12") can never create an invalid master row the
# HSN editor would itself reject.
_HSN_CODE_RE = re.compile(r"^[0-9]{2,12}$")

_MAX_NAME_LEN = 200
_MAX_CODE_LEN = 40
_MAX_SHORT_LEN = 120
_MAX_UOM_LEN = 20
_MAX_HSN_LEN = 10
# `md_hsn.description` column cap — a synced product name is truncated to fit.
_MAX_HSN_DESC_LEN = 300

# Required (non-null) string fields and their column caps.
_REQUIRED_FIELDS: dict[str, int] = {"name": _MAX_NAME_LEN, "uom": _MAX_UOM_LEN}
# Optional string fields (empty collapses to None) and their column caps.
_OPTIONAL_FIELDS: dict[str, int] = {
    "code": _MAX_CODE_LEN,
    "brand": _MAX_SHORT_LEN,
    "model_number": _MAX_SHORT_LEN,
    "category": _MAX_SHORT_LEN,
    "hsn": _MAX_HSN_LEN,
}
# Every mutable field a PATCH may touch (adds the boolean `active` and the `gst_rate`).
_PATCHABLE: frozenset[str] = (
    frozenset(_REQUIRED_FIELDS) | frozenset(_OPTIONAL_FIELDS) | {"active", "gst_rate"}
)


class ProductError(Exception):
    """Raised for an invalid product operation (route maps to 422)."""


class DuplicateProduct(ProductError):
    """The product identity or code already exists (route maps to 409)."""


class ProductNotFound(ProductError):
    """The referenced product is missing (route maps to 404)."""


def _escape_like(term: str) -> str:
    r"""Escape LIKE wildcards so a user-typed `%`/`_` matches literally (paired with
    `escape="\\"` on `.ilike()`). The backslash is escaped first so it stays the
    escape char; without this a literal `%` in `q` would match every row."""
    return term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _clean_required(value: str, *, field: str, max_len: int) -> str:
    """Strip a required string; reject empty or over-cap so a direct service caller
    can't persist junk or exceed the column width."""
    cleaned = value.strip()
    if not cleaned:
        raise ProductError(f"{field} is required")
    if len(cleaned) > max_len:
        raise ProductError(f"{field} must be at most {max_len} characters")
    return cleaned


def _reject_reserved_code(code: str | None) -> None:
    """Reject a USER-supplied code that matches the auto namespace `PRD-######`
    (case-insensitive). Reserving the shape guarantees a hand-entered code can
    never collide with a system-minted `PRD-{id:06d}`. No-op for None."""
    if code is not None and _RESERVED_CODE_RE.match(code):
        raise ProductError(
            "codes of the form PRD-###### are reserved for system-generated product codes"
        )


def _clean_optional(value: str | None, *, field: str, max_len: int) -> str | None:
    """Strip an optional string; empty/whitespace collapses to None so identity
    stays consistent (the index treats NULL brand/model as '')."""
    if value is None:
        return None
    cleaned = value.strip()
    if not cleaned:
        return None
    if len(cleaned) > max_len:
        raise ProductError(f"{field} must be at most {max_len} characters")
    return cleaned


def _clean_gst_rate(value: Decimal | None) -> Decimal | None:
    """Validate an optional GST rate. None/absent clears it. A present value must be a
    Decimal in 0..100 with at most 2 decimal places — mirrors the PO line's `tax_rate`
    guard (range first, then reject an over-scale value rather than silently rounding it
    to `Numeric(5,2)`). Raises `ProductError` (route -> 422) on a bad value."""
    if value is None:
        return None
    if not isinstance(value, Decimal):
        raise ProductError("gst_rate must be a number")
    if value < 0 or value > 100:
        raise ProductError("gst_rate must be between 0 and 100")
    # Numeric(5,2): reject > 2 decimal places instead of letting Postgres round silently.
    if value != value.quantize(Decimal("0.01")):
        raise ProductError("gst_rate supports at most 2 decimal places")
    return value


def _upsert_hsn_master(
    db: Session,
    *,
    hsn: str | None,
    gst_rate: Decimal | None,
    description: str | None,
    actor_uid: str | None,
    allow_update: bool,
) -> tuple[str, Decimal | None]:
    """Keep the challan HSN master (`md_hsn`) in sync from a product's (hsn, gst_rate).

    NO-OP unless BOTH `hsn` and `gst_rate` are present AND `hsn` is a valid statutory code
    (`_HSN_CODE_RE`). Otherwise:
      * no `HsnCode` for that `hsn`  -> CREATE it (rate = the product's, description = the
        product name truncated to fit, active=True), audited — this is how a NEW hsn gets
        registered, always allowed;
      * exists with a DIFFERENT rate and `allow_update` -> UPDATE it (old->new audited);
      * exists with a DIFFERENT rate and NOT `allow_update` -> "protected" (leave it): a
        routine product save must never silently rewrite a printed statutory rate — only the
        EXPLICIT sync (allow_update=True) reconciles an existing rate;
      * exists with the same rate     -> no-op.

    Returns `(outcome, old_rate)` where outcome is
    "created" | "updated" | "protected" | "noop". The `hsn` key is unique, so a concurrent
    create is caught at flush and folded into the existing-row branch (never a 500)."""
    if not hsn or gst_rate is None or not _HSN_CODE_RE.match(hsn):
        return "noop", None

    existing = db.execute(
        select(HsnCode).where(HsnCode.hsn == hsn)
    ).scalar_one_or_none()
    if existing is None:
        row = HsnCode(
            hsn=hsn,
            gst_rate=gst_rate,
            description=(description or "")[:_MAX_HSN_DESC_LEN],
            active=True,
            updated_by=actor_uid,
        )
        try:
            # `with` releases the savepoint on success; a concurrent create of the same
            # hsn trips the UNIQUE index at flush and lands in the except.
            with db.begin_nested():
                db.add(row)
                db.flush()
        except IntegrityError:
            # Another writer created this hsn first — re-read and fall into update below.
            existing = db.execute(
                select(HsnCode).where(HsnCode.hsn == hsn)
            ).scalar_one_or_none()
            if existing is None:  # pragma: no cover - re-raise an unexpected integrity error
                raise
        else:
            audit.log(
                db,
                action="masterdata.hsn.synced",
                actor_uid=actor_uid,
                entity="md_hsn",
                entity_id=str(row.id),
                detail={"hsn": hsn, "gst_rate": str(gst_rate), "source": "product",
                        "created": True},
            )
            return "created", None

    # `existing` is set (either originally or after losing the create race).
    if existing.gst_rate == gst_rate:  # Decimal equality is by value (18 == 18.00)
        return "noop", existing.gst_rate
    if not allow_update:
        # A routine product save NEVER rewrites a curated statutory rate — leave it for the
        # explicit Sync (which reports every change). A stray product typo can't corrupt the
        # rate that prints on future challans for this hsn.
        return "protected", existing.gst_rate
    old_rate = existing.gst_rate
    existing.gst_rate = gst_rate
    existing.updated_by = actor_uid
    audit.log(
        db,
        action="masterdata.hsn.synced",
        actor_uid=actor_uid,
        entity="md_hsn",
        entity_id=str(existing.id),
        detail={"hsn": hsn, "old_gst_rate": str(old_rate), "new_gst_rate": str(gst_rate),
                "source": "product"},
    )
    return "updated", old_rate


def create_product(
    db: Session,
    *,
    name: str,
    code: str | None = None,
    brand: str | None = None,
    model_number: str | None = None,
    category: str | None = None,
    uom: str = "PCS",
    hsn: str | None = None,
    gst_rate: Decimal | None = None,
    actor_uid: str | None = None,
) -> Product:
    """Create a product, deduped case-insensitively on (name, brand, model_number).

    Strings are normalized (strip; empty optionals -> None). A collision on the
    identity index OR the unique `code` trips the DB backstop at flush and is
    surfaced as `DuplicateProduct` (route -> 409). Audited `product.created`.
    """
    clean_name = _clean_required(name, field="name", max_len=_MAX_NAME_LEN)
    clean_uom = _clean_required(uom, field="uom", max_len=_MAX_UOM_LEN)
    clean_code = _clean_optional(code, field="code", max_len=_MAX_CODE_LEN)
    _reject_reserved_code(clean_code)
    clean_brand = _clean_optional(brand, field="brand", max_len=_MAX_SHORT_LEN)
    clean_model = _clean_optional(model_number, field="model_number", max_len=_MAX_SHORT_LEN)
    clean_category = _clean_optional(category, field="category", max_len=_MAX_SHORT_LEN)
    clean_hsn = _clean_optional(hsn, field="hsn", max_len=_MAX_HSN_LEN)
    clean_gst_rate = _clean_gst_rate(gst_rate)

    product = Product(
        name=clean_name,
        code=clean_code,
        brand=clean_brand,
        model_number=clean_model,
        category=clean_category,
        uom=clean_uom,
        hsn=clean_hsn,
        gst_rate=clean_gst_rate,
        created_by=actor_uid,
    )
    try:
        # `with` releases the savepoint on success (no leak); a concurrent create
        # of the same identity/code trips a UNIQUE index at flush and lands here.
        with db.begin_nested():
            db.add(product)
            db.flush()
    except IntegrityError as exc:
        raise DuplicateProduct("a product with this identity or code already exists") from exc

    # Auto-mint a stable code for a code-less product from its now-assigned PK. Collision
    # is near-impossible going forward (the reserved-namespace guard blocks a user code in
    # the auto shape, and PKs are unique) — but LEGACY data created before the guard could
    # hold a hand-entered `PRD-######` equal to some later id's auto code, so map that flush's
    # IntegrityError to DuplicateProduct too (a clean 409 single-create / per-row bulk error,
    # never an unmapped 500 that poisons a whole bulk batch).
    if product.code is None:
        try:
            with db.begin_nested():
                product.code = f"PRD-{product.id:06d}"
                db.flush()
        except IntegrityError as exc:
            raise DuplicateProduct(
                f"the auto-generated code {product.code} collides an existing product code"
            ) from exc

    audit.log(
        db,
        action="product.created",
        actor_uid=actor_uid,
        entity="product",
        entity_id=str(product.id),
        detail={
            "name": clean_name,
            "code": product.code,  # the final code (user-supplied or auto-minted)
            "brand": clean_brand,
            "model_number": clean_model,
            "category": clean_category,
            "uom": clean_uom,
            "hsn": clean_hsn,
            "gst_rate": str(clean_gst_rate) if clean_gst_rate is not None else None,
        },
    )

    # Auto-register a NEW hsn into the challan HSN master (a fresh product is active).
    # allow_update=False: a product save never rewrites an EXISTING curated rate.
    _upsert_hsn_master(
        db, hsn=clean_hsn, gst_rate=clean_gst_rate, description=clean_name,
        actor_uid=actor_uid, allow_update=False,
    )
    return product


def get_product(db: Session, product_id: int) -> Product:
    """Return the product with `product_id`, or raise `ProductNotFound`."""
    product = db.get(Product, product_id)
    if product is None:
        raise ProductNotFound(f"product {product_id} not found")
    return product


def list_products(
    db: Session,
    *,
    q: str | None = None,
    category: str | None = None,
    active: bool | None = None,
    limit: int = 50,
    offset: int = 0,
    project_id: int | None = None,
) -> list[Product]:
    """List/search products. `q` is a LIKE-escaped substring across name/brand/
    model_number; `category` and `active` are exact filters.

    Picker curation: when `project_id` is given AND `q` is empty/blank AND that project
    has at least one tagged product, this returns ONLY that project's TAGGED products
    (still honoring `active`/`category`), ordered by name — the "curated default" an empty
    search box shows on a project's PO form. A project with NO tags falls back to the full
    catalogue, so an untagged project's picker behaves exactly like the un-scoped catalogue
    (the curation only NARROWS a picker once the project is tagged — no rollout regression
    where every existing project would otherwise open an empty picker). In every other case
    (a non-blank `q`, or no `project_id`) behavior is unchanged: the FULL catalogue is
    searched (newest first), keeping every product reachable by typing.
    """
    curated = project_id is not None and not (q and q.strip())
    if curated:
        # Only curate when the project actually has tags — else fall through to the full
        # catalogue (untagged project == plain picker; no empty-dropdown rollout cliff).
        has_tags = (
            db.execute(
                select(ProjectProduct.id)
                .where(ProjectProduct.project_id == project_id)
                .limit(1)
            ).first()
            is not None
        )
        curated = has_tags
    if curated:
        stmt = (
            select(Product)
            .join(ProjectProduct, ProjectProduct.product_id == Product.id)
            .where(ProjectProduct.project_id == project_id)
            .order_by(Product.name)
        )
    else:
        stmt = select(Product).order_by(Product.id.desc())
        if q:
            like = f"%{_escape_like(q.strip())}%"
            stmt = stmt.where(
                Product.name.ilike(like, escape="\\")
                | Product.brand.ilike(like, escape="\\")
                | Product.model_number.ilike(like, escape="\\")
            )
    if category:
        stmt = stmt.where(func.lower(Product.category) == category.strip().lower())
    if active is not None:
        stmt = stmt.where(Product.active == active)
    stmt = stmt.limit(limit).offset(offset)
    return list(db.execute(stmt).scalars())


def update_product(
    db: Session,
    *,
    product_id: int,
    fields: dict[str, Any],
    actor_uid: str | None = None,
) -> Product:
    """Patch a product with only the keys present in `fields` (the route passes the
    client's `exclude_unset` dump, so an omitted key is left untouched while an
    explicit `null` clears an optional field).

    A change that collides with another product's identity/code trips the DB
    backstop and is surfaced as `DuplicateProduct` (route -> 409). Missing row
    raises `ProductNotFound` (route -> 404). Audited `product.updated`.
    """
    unknown = set(fields) - _PATCHABLE
    if unknown:
        raise ProductError(f"unknown field(s): {sorted(unknown)}")

    product = get_product(db, product_id)
    changed: dict[str, Any] = {}

    for field, max_len in _REQUIRED_FIELDS.items():
        if field in fields:
            value = fields[field]
            if not isinstance(value, str):
                raise ProductError(f"{field} must be a string")
            required_clean = _clean_required(value, field=field, max_len=max_len)
            setattr(product, field, required_clean)
            changed[field] = required_clean

    for field, max_len in _OPTIONAL_FIELDS.items():
        if field in fields:
            value = fields[field]
            if value is not None and not isinstance(value, str):
                raise ProductError(f"{field} must be a string or null")
            optional_clean = _clean_optional(value, field=field, max_len=max_len)
            if field == "code":
                _reject_reserved_code(optional_clean)
            setattr(product, field, optional_clean)
            changed[field] = optional_clean

    if "active" in fields:
        value = fields["active"]
        if not isinstance(value, bool):
            raise ProductError("active must be a boolean")
        product.active = value
        changed["active"] = value

    if "gst_rate" in fields:
        # bool is a subclass of int (not Decimal); an explicit null clears the rate.
        clean_rate = _clean_gst_rate(fields["gst_rate"])
        product.gst_rate = clean_rate
        changed["gst_rate"] = str(clean_rate) if clean_rate is not None else None

    try:
        with db.begin_nested():
            db.flush()  # trips the identity/code unique backstop if this edit collides
    except IntegrityError as exc:
        raise DuplicateProduct("a product with this identity or code already exists") from exc

    audit.log(
        db,
        action="product.updated",
        actor_uid=actor_uid,
        entity="product",
        entity_id=str(product.id),
        detail=changed,
    )

    # Sync the challan HSN master from the product's CURRENT (hsn, gst_rate) — but ONLY for an
    # ACTIVE product (a soft-deleted product must never drive the live statutory rate), and
    # allow_update=False so a routine edit auto-registers a new hsn but never rewrites an
    # existing curated rate. No-op unless both hsn & gst_rate are present.
    if product.active:
        _upsert_hsn_master(
            db, hsn=product.hsn, gst_rate=product.gst_rate, description=product.name,
            actor_uid=actor_uid, allow_update=False,
        )
    return product


# ------------------------------------------------------------- bulk (Excel)
#
# CREATE/UPDATE the Product Master from an .xlsx (mirrors the PO bulk-upload). Every row is
# resolved to an existing product (a MATCH -> update) or a new one (NO match -> create), and
# each row runs inside its OWN SAVEPOINT so one bad row can never poison the rest of the batch.
#
# Match rule (per row):
#   * a `code` cell present -> match that exact product code (the unique `code` key);
#   * else -> match by IDENTITY = case-insensitive (name, brand, model_number), exactly the
#     tuple the `uq_product_identity` DB index normalizes (lower(name),
#     lower(coalesce(brand,'')), lower(coalesce(model_number,''))).
# A match UPDATES the product with the row's provided (non-empty) fields; no match CREATEs it
# (code auto-minted when the row omits one). `create_product`/`update_product` are reused
# VERBATIM, so their identity/code-collision + reserved-code (`PRD-######`) guards apply and any
# `ProductError` surfaces as a per-row error rather than crashing the batch.

# Canonical bulk columns -> the friendly headers accepted for each (case-insensitive,
# whitespace/hyphen collapsed to '_'). Only `name` is required.
_BULK_ALIASES: dict[str, str] = {
    "name": "name", "product_name": "name", "product": "name",
    "item": "name", "item_name": "name",
    "code": "code", "product_code": "code", "sku": "code", "item_code": "code",
    "brand": "brand", "make": "brand", "manufacturer": "brand",
    "model_number": "model_number", "model": "model_number", "model_no": "model_number",
    "category": "category", "cat": "category",
    "uom": "uom", "unit": "uom", "units": "uom", "unit_of_measure": "uom",
    "hsn": "hsn", "hsn_code": "hsn",
    "gst_rate": "gst_rate", "gst": "gst_rate", "gst_percent": "gst_rate", "tax": "gst_rate",
}

# The downloadable TEMPLATE: the canonical header in a fixed, friendly order + one example row.
# `code` is left blank so the example creates (auto-minting a code); the operator fills a code
# only to set a custom SKU or to update-by-code.
_BULK_TEMPLATE_COLUMNS: tuple[str, ...] = (
    "name", "code", "brand", "model_number", "category", "uom", "hsn", "gst_rate",
)
_BULK_TEMPLATE_EXAMPLE: tuple[object, ...] = (
    "Sample Product — replace with your own", "", "Acme", "GX-100", "Appliances", "PCS",
    "8509", 18,
)


@dataclass
class ProductsBulkResult:
    """Outcome of an Excel product upsert: the CODES created / updated, and per-row errors
    (1-based sheet row + message). A row is in exactly one bucket."""

    created: list[str] = dc_field(default_factory=list)
    updated: list[str] = dc_field(default_factory=list)
    errors: list[tuple[int, str]] = dc_field(default_factory=list)


def build_products_template_xlsx() -> bytes:
    """Return the bulk-upsert template as ``.xlsx`` bytes: the header row in the exact column
    order the parser accepts + one illustrative example row."""
    wb = Workbook()
    ws = wb.active or wb.create_sheet()
    ws.title = "products"
    ws.append(list(_BULK_TEMPLATE_COLUMNS))
    ws.append(list(_BULK_TEMPLATE_EXAMPLE))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm_header(raw: str) -> str:
    return "_".join(raw.strip().lower().replace("-", " ").split())


def _cell_str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()


def _parse_products_workbook(
    data: bytes,
) -> tuple[list[tuple[int, dict[str, str]]], list[tuple[int, str]]]:
    """Read the first worksheet against the bulk template. Returns (rows, errors); a bad file
    or a missing required ``name`` header yields ([], [errors])."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises many types on a bad file
        return [], [(1, "could not read the file as an .xlsx workbook")]
    try:
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        col_index: dict[str, int] = {}
        for idx, raw in enumerate(header or ()):
            if raw is None:
                continue
            key = _BULK_ALIASES.get(_norm_header(str(raw)))
            if key is not None and key not in col_index:
                col_index[key] = idx

        if "name" not in col_index:
            return [], [(1, "required column 'name' is missing")]

        rows: list[tuple[int, dict[str, str]]] = []
        for sheet_row, values in enumerate(rows_iter, start=2):
            cells = {
                key: _cell_str(values[idx] if idx < len(values) else None)
                for key, idx in col_index.items()
            }
            if all(v == "" for v in cells.values()):
                continue  # skip a fully-blank row
            rows.append((sheet_row, cells))
        return rows, []
    finally:
        wb.close()


def _match_bulk_product(
    db: Session, *, name: str, code: str, brand: str, model: str
) -> Product | None:
    """Resolve a row to an existing product: by exact ``code`` when one is supplied, else by
    the case-insensitive identity tuple the unique index enforces. Returns None (-> create)
    when nothing matches. Both keys are unique, so at most one row can match."""
    if code:
        return db.execute(
            select(Product).where(Product.code == code)
        ).scalar_one_or_none()
    return db.execute(
        select(Product).where(
            func.lower(Product.name) == name.lower(),
            func.lower(func.coalesce(Product.brand, "")) == brand.lower(),
            func.lower(func.coalesce(Product.model_number, "")) == model.lower(),
        )
    ).scalar_one_or_none()


def _bulk_update_fields(cells: dict[str, str]) -> dict[str, Any]:
    """The provided (non-empty) descriptive fields for an UPDATE. A blank cell is left
    untouched (bulk never CLEARS a field — unlike the PATCH API's explicit-null semantics)."""
    fields: dict[str, Any] = {}
    for col in ("name", "code", "brand", "model_number", "category", "uom", "hsn"):
        value = cells.get(col, "").strip()
        if value:
            fields[col] = value
    return fields


def _bulk_gst_rate(cells: dict[str, str]) -> Decimal | None:
    """Parse the optional `gst_rate` cell to a Decimal (None when blank). A non-numeric cell
    is a ``ProductError`` (per-row error). Range/scale (0..100, <=2 dp) is enforced downstream
    by `create_product`/`update_product` via `_clean_gst_rate` — the SAME guard as the API."""
    raw = cells.get("gst_rate", "").strip()
    if not raw:
        return None
    parsed = parse_qty(raw)
    if parsed is None:
        raise ProductError("gst_rate must be a number between 0 and 100")
    return parsed


def _apply_bulk_row(
    db: Session, cells: dict[str, str], *, actor_uid: str | None
) -> tuple[str, str]:
    """Create or update one product from a row. Returns ("created"|"updated", code). Raises
    ``ProductError`` (incl. ``DuplicateProduct``) for a bad/colliding/reserved-code row."""
    name = cells.get("name", "").strip()
    if not name:
        raise ProductError("name is required")
    code = cells.get("code", "").strip()
    brand = cells.get("brand", "").strip()
    model = cells.get("model_number", "").strip()
    gst_rate = _bulk_gst_rate(cells)  # non-numeric -> ProductError (per-row)

    match = _match_bulk_product(db, name=name, code=code, brand=brand, model=model)
    if match is None:
        product = create_product(
            db,
            name=name,
            code=code or None,
            brand=brand or None,
            model_number=model or None,
            category=cells.get("category", "").strip() or None,
            uom=cells.get("uom", "").strip() or "PCS",
            hsn=cells.get("hsn", "").strip() or None,
            gst_rate=gst_rate,
            actor_uid=actor_uid,
        )
        return "created", product.code or ""

    fields = _bulk_update_fields(cells)
    if cells.get("gst_rate", "").strip():
        # Only inject when the cell carries a value — a blank cell leaves the rate untouched
        # (bulk never CLEARS a field), matching the other descriptive columns.
        fields["gst_rate"] = gst_rate
    product = update_product(db, product_id=match.id, fields=fields, actor_uid=actor_uid)
    return "updated", product.code or ""


def bulk_upsert_from_excel(
    db: Session, data: bytes, *, actor_uid: str | None = None
) -> ProductsBulkResult:
    """Parse an ``.xlsx`` and CREATE/UPDATE one product per row. NO row is silently dropped: a
    malformed / colliding / reserved-code row lands in ``errors`` (with its 1-based sheet row)
    while every other row still processes — each row runs inside its OWN SAVEPOINT so a failure
    rolls back only that row, never the batch. Returns created / updated codes + errors. Caller
    commits."""
    rows, parse_errors = _parse_products_workbook(data)
    result = ProductsBulkResult(errors=list(parse_errors))
    if parse_errors:  # header / file problem — nothing to upsert
        return result

    for sheet_row, cells in rows:
        try:
            # A per-row SAVEPOINT: any error (validation, identity/code collision, reserved
            # code) rolls back ONLY this row's writes so the surviving rows still commit.
            with db.begin_nested():
                outcome, code = _apply_bulk_row(db, cells, actor_uid=actor_uid)
        except ProductError as exc:  # DuplicateProduct is a ProductError subclass
            result.errors.append((sheet_row, str(exc)))
            continue
        (result.created if outcome == "created" else result.updated).append(code)

    return result


# ------------------------------------------------- manual HSN-master reconciliation
#
# A one-shot "sync all" that rebuilds `md_hsn` from every product carrying BOTH an `hsn`
# and a `gst_rate`. Single-record create/update already sync as they go; this is the manual
# backfill/repair an operator runs after entering rates in bulk.


@dataclass
class HsnSyncResult:
    """Outcome of a full product -> `md_hsn` reconciliation.

    * ``created``   — HSN codes newly inserted into `md_hsn`.
    * ``updated``   — HSN codes whose rate was changed, each ``{hsn, old_rate, new_rate}``.
    * ``conflicts`` — HSNs where >=2 ACTIVE products disagree on the rate, each
      ``{hsn, rates, product_ids}``. A conflict is still APPLIED (deterministically, using
      the most-recently-created active product's rate) so the run is STABLE; the bucket
      flags the data for an operator to fix.
    """

    created: list[str] = dc_field(default_factory=list)
    updated: list[dict[str, str]] = dc_field(default_factory=list)
    conflicts: list[dict[str, Any]] = dc_field(default_factory=list)


def _rate_winner(products: list[Product]) -> Product:
    """The most-recently-created product (created_at, then id as a stable tiebreak) — the
    deterministic choice whose rate wins for its hsn."""
    return max(products, key=lambda p: (p.created_at, p.id))


def sync_hsn_master_from_products(
    db: Session, *, actor_uid: str | None = None
) -> HsnSyncResult:
    """Reconcile `md_hsn` from every ACTIVE product with BOTH `hsn` and `gst_rate` set.

    Products are grouped by `hsn`; a group with NO active product is skipped entirely (a
    soft-deleted product must not drive the live statutory rate). The authoritative rate is
    the most-recently-created ACTIVE product's; when two or more active products carry
    DIFFERENT rates the group is reported as a conflict AND still applied deterministically
    (so re-running never thrashes `md_hsn`). Unlike a routine product save, this explicit
    reconcile passes `allow_update=True`, so it MAY change an existing curated rate — that is
    the deliberate, operator-initiated path, and every change is reported + audited. Caller
    commits."""
    result = HsnSyncResult()
    products = list(
        db.execute(
            select(Product)
            .where(Product.hsn.is_not(None), Product.gst_rate.is_not(None))
            .order_by(Product.id)
        ).scalars()
    )

    groups: dict[str, list[Product]] = {}
    for product in products:
        if product.hsn is None or product.gst_rate is None:  # narrows for the type checker
            continue
        groups.setdefault(product.hsn, []).append(product)

    for hsn, group in groups.items():
        active = [p for p in group if p.active]
        if not active:
            # No ACTIVE product carries this hsn — a soft-deleted product must not drive the
            # live statutory rate, so skip it entirely (never seed/change md_hsn from it).
            continue
        # Distinct rates BY VALUE among active products (Decimal 18 == 18.00).
        distinct_active = {p.gst_rate for p in active}
        if len(distinct_active) >= 2:
            result.conflicts.append({
                "hsn": hsn,
                "rates": sorted(str(r) for r in distinct_active),
                "product_ids": sorted(p.id for p in active),
            })
        winner = _rate_winner(active)
        # Re-validate the winner's rate — a rate written by a direct DB import bypasses the
        # route/bulk guards; never propagate an out-of-range rate to the statutory master.
        winner_rate = _clean_gst_rate(winner.gst_rate)

        outcome, old_rate = _upsert_hsn_master(
            db, hsn=hsn, gst_rate=winner_rate, description=winner.name,
            actor_uid=actor_uid, allow_update=True,
        )
        if outcome == "created":
            result.created.append(hsn)
        elif outcome == "updated":
            result.updated.append(
                {"hsn": hsn, "old_rate": str(old_rate), "new_rate": str(winner_rate)}
            )

    return result
