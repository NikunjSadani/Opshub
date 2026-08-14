"""Self-documenting Excel upload template for the challan generator.

Builds a two-sheet `.xlsx`:

* **Challans** — the data sheet the operator fills in (this MUST be the first
  worksheet, because the parser reads `worksheets[0]`). Row 1 is the FRIENDLY
  header (`schema.COLUMN_HEADERS`, in `schema.CHALLAN_COLUMNS` order), styled +
  frozen, with a hover-comment on every header cell explaining that column.
  Worked example rows follow (two challans; the second is multi-line).
* **Instructions** — a full reference: every column (Required?/Applies to/what to
  enter, all driven off the schema so they can't drift) plus Dos & Don'ts.

The friendly headers round-trip: the parser matches each one (and the raw keys
and the alias set) via `schema.header_to_key`, so a downloaded template uploads
back cleanly.
"""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from app.modules.challan.schema import (
    CHALLAN_COLUMNS,
    COLUMN_HEADERS,
    GROUP_CONSISTENT_FIELDS,
    REQUIRED_COLUMNS,
)

# Per-column prose (Required?/Applies-to are DERIVED from the schema below).
_HELP: dict[str, str] = {
    "challan_group": "A label that groups rows into ONE challan. Rows sharing the "
                     "same value become a single (multi-line) challan, e.g. C1, C2.",
    "project_id": "The Project this challan is for, e.g. BRI-001. It must ALREADY "
                  "EXIST and be Active in the Projects module. Printed on the challan.",
    "ship_to_enterprise": "Enterprise / business name of the ship-to party "
                          "(the 'Detail of Shipment to' block).",
    "ship_to_name": "Name of the party the goods are shipped to.",
    "ship_to_address_line1": "Ship-to address, line 1. The address is captured in "
                             "parts (line 1/2, city, state, pincode) and printed as "
                             "one block on the challan.",
    "ship_to_address_line2": "Ship-to address, line 2 (optional continuation of the "
                             "street address).",
    "ship_to_city": "Ship-to city / town.",
    "ship_to_state": "Ship-to (destination) state.",
    "ship_to_pincode": "Ship-to PIN code. Optional; if given it must be exactly "
                       "6 digits.",
    "ship_to_phone": "Contact phone number at the ship-to party.",
    "consignee_name": "Name of the consignee (the bill-to party). Typed inline and "
                      "matched to the party master by its GSTIN.",
    "consignee_address_line1": "Consignee address, line 1. Printed as one joined "
                               "block on the challan.",
    "consignee_address_line2": "Consignee address, line 2 (optional continuation).",
    "consignee_pincode": "Consignee PIN code. Optional; if given it must be exactly "
                         "6 digits.",
    "consignee_state": "Consignee state.",
    "consignee_phone": "Contact phone number at the consignee.",
    "consignee_gstin": "The consignee's 15-character GSTIN — the key the party is "
                       "identified by. A new GSTIN is auto-added to the party master; "
                       "a known GSTIN with different details is flagged as a deviation "
                       "and the STORED record is used.",
    "challan_date": "Challan date. Formats: DD-MM-YYYY (e.g. 15-05-2026), "
                    "YYYY-MM-DD, or a real Excel date. This sets the financial "
                    "year (Apr-Mar) in the challan number.",
    "description": "Description of the line item (product).",
    "hsn": "HSN code of the line item. Must already exist as an active HSN in "
           "Master Data (it supplies the GST rate).",
    "quantity": "Quantity for the line. Must be a number greater than 0 "
                "(decimals allowed).",
    "rate": "Pre-tax unit rate. Optional. May be free text like 'as per "
            "contract'. If a number, it must not be negative.",
    "amount": "Tax-INCLUSIVE line amount in rupees = Rate x Qty x (1 + GST%). "
              "Optional — leave blank for a value-free challan. Within one "
              "challan, either ALL lines have an amount or NONE do.",
    "gst_rate": "GST % for the line. Optional. If given, it must equal the HSN's "
                "configured rate (a safety cross-check).",
    "po_number": "Purchase-order number. Optional. Stored but NOT printed on the "
                 "challan.",
    "invoice_number": "Invoice number. Optional. Printed just above the challan "
                      "number.",
}

# Worked examples (keys are the canonical schema keys, not the friendly headers):
# C1 is a single priced line (HSN 1509 @ 5% -> 100 x 10 x 1.05 = 1050.00); C2 is a
# two-line value-free challan to a different destination + consignee.
_EXAMPLES: tuple[dict[str, str], ...] = (
    {
        "challan_group": "C1", "project_id": "BRI-001",
        "ship_to_enterprise": "Acme Enterprises",
        "ship_to_name": "Acme Retail Pvt Ltd",
        "ship_to_address_line1": "12 MG Road", "ship_to_address_line2": "Fort",
        "ship_to_city": "Mumbai", "ship_to_state": "Maharashtra",
        "ship_to_pincode": "400001", "ship_to_phone": "9900000000",
        "consignee_name": "Umang Foods LLP",
        "consignee_address_line1": "45 Industrial Estate",
        "consignee_address_line2": "Andheri East",
        "consignee_pincode": "400069", "consignee_state": "Maharashtra",
        "consignee_phone": "9900011111", "consignee_gstin": "27AAPFU0939F1ZV",
        "challan_date": "15-05-2026", "description": "Olive Oil 1L", "hsn": "1509",
        "quantity": "10", "rate": "100.00", "amount": "1050.00", "gst_rate": "5",
        "po_number": "PO-2026-100", "invoice_number": "INV-2026-100",
    },
    {
        "challan_group": "C2", "project_id": "BRI-002",
        "ship_to_enterprise": "Best Foods LLP",
        "ship_to_name": "Best Foods",
        "ship_to_address_line1": "9 Link Road", "ship_to_address_line2": "Model Colony",
        "ship_to_city": "Pune", "ship_to_state": "Maharashtra",
        "ship_to_pincode": "411001", "ship_to_phone": "9800000000",
        "consignee_name": "Bharat Beverages Pvt Ltd",
        "consignee_address_line1": "22 GIDC", "consignee_address_line2": "Vatva",
        "consignee_pincode": "382445", "consignee_state": "Gujarat",
        "consignee_phone": "9800011111", "consignee_gstin": "24AAACB2894G1ZT",
        "challan_date": "16-05-2026", "description": "Sample Pack A", "hsn": "1509",
        "quantity": "2", "rate": "", "amount": "", "gst_rate": "",
        "po_number": "", "invoice_number": "",
    },
    {
        "challan_group": "C2", "project_id": "BRI-002",
        "ship_to_enterprise": "Best Foods LLP",
        "ship_to_name": "Best Foods",
        "ship_to_address_line1": "9 Link Road", "ship_to_address_line2": "Model Colony",
        "ship_to_city": "Pune", "ship_to_state": "Maharashtra",
        "ship_to_pincode": "411001", "ship_to_phone": "9800000000",
        "consignee_name": "Bharat Beverages Pvt Ltd",
        "consignee_address_line1": "22 GIDC", "consignee_address_line2": "Vatva",
        "consignee_pincode": "382445", "consignee_state": "Gujarat",
        "consignee_phone": "9800011111", "consignee_gstin": "24AAACB2894G1ZT",
        "challan_date": "16-05-2026", "description": "Sample Pack B", "hsn": "1509",
        "quantity": "3", "rate": "", "amount": "", "gst_rate": "",
        "po_number": "", "invoice_number": "",
    },
)

_DOS: tuple[str, ...] = (
    "Repeat the Challan Group value to add more lines to one challan; keep every "
    "challan-level field identical across those rows.",
    "Fill every Required column (see the table above).",
    "Use a Project ID that already exists and is Active in the Projects module.",
    "The consignee is identified by its GSTIN — a new GSTIN is added automatically.",
    "Enter a Quantity greater than 0.",
    "If you enter an Amount, use the tax-inclusive figure = Rate x Qty x (1 + GST%).",
    "Replace the example rows on the 'Challans' sheet with your own data.",
    "Save the file as .xlsx before uploading.",
)

_DONTS: tuple[str, ...] = (
    "Don't rename, remove, reorder, or add header columns — keep row 1 exactly as "
    "provided.",
    "Don't reuse one Challan Group for shipments to different destinations.",
    "Don't invent a Project ID — create it in Projects first.",
    "Don't mix priced and value-free lines in the same challan (all priced, or all "
    "blank).",
    "Don't enter a GST rate different from the HSN's configured rate.",
    "Don't upload a .csv — the uploader accepts .xlsx only.",
)

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")  # slate-800
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_REQ_FILL = PatternFill("solid", fgColor="FEF3C7")      # amber-100 (required cols)
_TITLE_FONT = Font(bold=True, size=14)
_SECTION_FONT = Font(bold=True, size=11, color="1F2937")
_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_THIN = Side(style="thin", color="D1D5DB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _applies_to(col: str) -> str:
    if col == "challan_group":
        return "Grouping key"
    if col in GROUP_CONSISTENT_FIELDS:
        return "Whole challan"
    return "Per line"


def _build_data_sheet(ws: Worksheet) -> None:
    ws.title = "Challans"
    ws.freeze_panes = "A2"
    for idx, col in enumerate(CHALLAN_COLUMNS, start=1):
        header = COLUMN_HEADERS[col]
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill = _HEADER_FILL if col not in REQUIRED_COLUMNS else _REQ_FILL
        cell.font = _HEADER_FONT if col not in REQUIRED_COLUMNS else Font(bold=True)
        req = "REQUIRED" if col in REQUIRED_COLUMNS else "Optional"
        cell.comment = Comment(f"{req} · {_applies_to(col)}\n\n{_HELP[col]}", "OpsHub")
        ws.column_dimensions[cell.column_letter].width = max(14, len(header) + 4)
    for r, example in enumerate(_EXAMPLES, start=2):
        for idx, col in enumerate(CHALLAN_COLUMNS, start=1):
            ws.cell(row=r, column=idx, value=example.get(col, ""))


def _build_instructions_sheet(ws: Worksheet) -> None:
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 80

    ws["A1"] = "Delivery Challan — upload template"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = (
        "Fill the 'Challans' sheet (one row per line item) and upload it as .xlsx. "
        "The consignor is fixed; the consignee is identified by its GSTIN; the "
        "Project ID must already exist in Projects."
    )
    ws["A2"].alignment = _WRAP_TOP
    ws.merge_cells("A2:D2")

    head = 4
    for i, label in enumerate(("Column", "Required?", "Applies to", "What to enter")):
        c = ws.cell(row=head, column=1 + i, value=label)
        c.font = _HEADER_FONT
        c.fill = _HEADER_FILL
        c.border = _BORDER
    row = head + 1
    for col in CHALLAN_COLUMNS:
        req = "Required" if col in REQUIRED_COLUMNS else "Optional"
        values = (COLUMN_HEADERS[col], req, _applies_to(col), _HELP[col])
        for i, val in enumerate(values):
            c = ws.cell(row=row, column=1 + i, value=val)
            c.alignment = _WRAP_TOP
            c.border = _BORDER
            if i == 0:
                c.font = Font(bold=True)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="DOs").font = _SECTION_FONT
    row += 1
    for item in _DOS:
        cell = ws.cell(row=row, column=1, value=f"✓  {item}")
        cell.alignment = _WRAP_TOP
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="DON'Ts").font = _SECTION_FONT
    row += 1
    for item in _DONTS:
        cell = ws.cell(row=row, column=1, value=f"✗  {item}")
        cell.alignment = _WRAP_TOP
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        row += 1


def build_template_xlsx() -> bytes:
    """Return the self-documenting upload template as `.xlsx` bytes."""
    wb = Workbook()
    _build_data_sheet(wb.active)  # the default sheet becomes 'Challans' (worksheets[0])
    _build_instructions_sheet(wb.create_sheet())
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
