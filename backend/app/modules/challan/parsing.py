"""Challan Excel parsing + structural validation — the no-database front door.

This layer turns an uploaded `.xlsx` into `RawRow`s and reports every problem it
can find WITHOUT touching master data, so a bad spreadsheet is rejected with a
complete English report before the service ever reserves a number:

  parse_workbook(bytes)          -> (list[RawRow], list[RowError])  # shape/headers
  structural_row_errors(rows)    -> list[RowError]                  # per-row + group

The 26-column template captures a structured ship-to block (line1/line2/city/
pincode/state), an INLINE consignee (name / address / GSTIN — the golden-record
key), a referenced `project_id`, and a line item, with rows sharing an explicit
`challan_group` folded into one challan. Uploaded headers are matched via
`schema.header_to_key`, so the friendly headers ("Project ID", "Consignee
GSTIN", "GST No", "City", ...), the raw keys, and the alias set all parse.

The semantic layer (consignor lookup, GSTIN checksum + golden-record resolution,
HSN lookups, building `ParsedChallan`s) lives in the service — this module
deliberately stops at structural truth (no GSTIN checksum here).

Money is integer PAISE via `Decimal`, never float: `parse_paise` multiplies rupees
by 100 and rounds HALF-UP to the paise. `rate` may legitimately be free text
(e.g. "as per contract"), so `parse_rate` returns the verbatim text plus an
optional parsed paise value.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from openpyxl import load_workbook

from app.modules.challan.schema import (
    GROUP_CONSISTENT_FIELDS,
    REQUIRED_COLUMNS,
    RawRow,
    RowError,
    header_to_key,
)

# Excel's day-zero for the 1900 date system, offset by its fictional 1900-02-29.
_EXCEL_EPOCH = date(1899, 12, 30)
# Excel's own serial ceiling: 2958465 == 9999-12-31. Anything larger overflows a
# Python `date`, so we reject it as "not a date" rather than let it raise.
_MAX_EXCEL_SERIAL = Decimal(2_958_465)
_DATE_FORMATS = ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y")


# ------------------------------------------------------------------- helpers

# Only plain ASCII numerics are a number here. This rejects non-ASCII digit code
# points (Devanagari/Arabic-Indic/fullwidth) and `Decimal`'s own "1_000"
# underscore grouping — both of which `Decimal()` would otherwise silently accept
# — mirroring the pincode ASCII guard so they surface as clean row errors. It also
# incidentally rejects "NaN"/"Infinity"/"1E30" (letters), which must never
# masquerade as a quantity or amount.
_ASCII_NUMERIC = re.compile(r"\A[+\-0-9.,]+\Z")


def _degroup(s: str) -> str | None:
    """Validate thousands-comma grouping and return the comma-free number, else None.

    A plain number with no comma passes through unchanged. When commas ARE present
    they must form valid grouping — Western `1,234,567.89` or Indian `1,23,456.78`,
    both accepted — meaning: the rightmost integer group is exactly 3 digits, the
    first group is 1-3 digits, and every interior group has a consistent width (all
    3 = Western, all 2 = Indian). A comma AFTER the decimal point (a European decimal
    comma, e.g. "1.234,50"), an empty/oversized group, or mixed interior widths
    (e.g. "1,0,0") is malformed grouping -> None, so it becomes a clean validation
    row error instead of a silently-wrong value.
    """
    if "," not in s:
        return s
    sign, body = ("", s)
    if body[:1] in "+-":
        sign, body = body[0], body[1:]
    if "." in body:
        int_part, _, frac_part = body.partition(".")
        if "," in frac_part:  # a comma may never trail the decimal point
            return None
    else:
        int_part, frac_part = body, None
    groups = int_part.split(",")
    if len(groups) < 2 or any(not g.isdigit() for g in groups):
        return None
    first, last, interior = groups[0], groups[-1], groups[1:-1]
    if not (1 <= len(first) <= 3) or len(last) != 3:
        return None
    if interior and not (all(len(g) == 3 for g in interior)
                         or all(len(g) == 2 for g in interior)):
        return None
    degrouped = sign + "".join(groups)
    return degrouped if frac_part is None else f"{degrouped}.{frac_part}"


def _to_decimal(text: str) -> Decimal | None:
    """Parse a (comma-grouped) numeric string to a finite Decimal, else None.

    Rejects blanks, junk, non-ASCII digits, malformed comma grouping, and
    non-finite tokens — anything that isn't an unambiguous plain-ASCII number
    becomes None so it surfaces as a clean row error rather than a wrong value.
    """
    s = text.strip()
    if not s:
        return None
    if not _ASCII_NUMERIC.match(s):  # L1: non-ASCII digits / "1_000" underscores
        return None
    degrouped = _degroup(s)  # M2: reject European "1.234,50" / ambiguous "1,0,0"
    if degrouped is None:
        return None
    try:
        value = Decimal(degrouped)
    except InvalidOperation:
        return None
    return value if value.is_finite() else None


def parse_paise(text: str) -> int | None:
    """Rupees string ("1,234.50") -> integer paise (123450), or None if not numeric.

    Half-up rounding to the paise keeps money exact and float-free. A finite but
    absurdly-large value (e.g. a 27-digit rupee figure) overflows the Decimal
    context on quantize, raising InvalidOperation; we catch it and return None so a
    junk cell becomes a clean row error instead of crashing the synchronous upload
    with a 500. (A shorter "1E30" is already rejected upstream as non-numeric.)
    """
    rupees = _to_decimal(text)
    if rupees is None:
        return None
    try:
        paise = (rupees * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None
    return int(paise)


def parse_rate(text: str) -> tuple[str, int | None]:
    """Return (verbatim text, paise) — rate may be free text, so paise is optional."""
    return text.strip(), parse_paise(text)


def parse_qty(text: str) -> Decimal | None:
    """Parse a quantity/rate figure to a finite Decimal, or None if not numeric."""
    return _to_decimal(text)


def parse_date(text: str) -> date | None:
    """Parse DD-MM-YYYY, YYYY-MM-DD, DD/MM/YYYY, an ISO datetime, or an Excel serial.

    openpyxl hands date cells back as `datetime`s that we stringify to ISO, so the
    ISO/`datetime`-string branch covers cells authored as real Excel dates; the
    serial branch covers a bare numeric date. Returns None if nothing matches.
    """
    s = text.strip()
    if not s:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    try:
        return date.fromisoformat(s[:10])  # "2026-08-10" or "2026-08-10 00:00:00"
    except ValueError:
        pass
    serial = _to_decimal(s)
    # Bound the serial to Excel's own date range (1..2958465 == 9999-12-31); a junk
    # numeric like "9999999" would otherwise overflow `date` and raise, escaping the
    # row-error path as a 500. Out-of-range -> a normal "not a date" row error.
    if serial is None or serial <= 0 or serial > _MAX_EXCEL_SERIAL:
        return None
    try:
        return _EXCEL_EPOCH + timedelta(days=int(serial))
    except (OverflowError, ValueError):  # pragma: no cover - bound above already guards
        return None


# XML-illegal control characters (all C0 except tab/newline/carriage-return). openpyxl
# raises IllegalCharacterError writing these, so a cell carrying one would 500 the
# downstream Excel review report — strip them at coercion so every path stays clean.
_ILLEGAL_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _to_str(value: object) -> str:
    """Coerce a raw cell value to a trimmed string ("" for blank), stripping
    XML-illegal control characters from free text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        # Avoid "10.0" for whole-number Excel cells; keep exact decimals otherwise.
        return str(int(value)) if value.is_integer() else str(value)
    if isinstance(value, int):
        return str(value)
    return _ILLEGAL_CTRL.sub("", str(value)).strip()


# ------------------------------------------------------------------- parse

def parse_workbook(data: bytes) -> tuple[list[RawRow], list[RowError]]:
    """Load the first worksheet from `data` and read its rows against the template.

    Row 1 is the header; each header cell is mapped to its canonical column key
    via `schema.header_to_key`, which accepts the friendly display header, the
    raw key, and the alias set (all case-insensitive + whitespace-collapsed).
    Returns ([], [errors]) — never raises — if the file can't be opened or any
    `REQUIRED_COLUMNS` header is missing (one error per missing key, all at row
    1). Otherwise returns (rows, []): each `RawRow` carries the real 1-based
    sheet row number (first data row = 2), every mapped cell coerced to a
    trimmed string; fully-blank rows and unknown/extra columns are dropped.
    """
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises many types on a bad file
        return [], [RowError(1, "", "could not read the file as an .xlsx workbook")]
    try:
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)

        col_index: dict[str, int] = {}
        for idx, raw in enumerate(header or ()):
            if raw is None:
                continue
            key = header_to_key(str(raw))
            if key is not None and key not in col_index:
                col_index[key] = idx

        missing = [c for c in REQUIRED_COLUMNS if c not in col_index]
        if missing:
            return [], [
                RowError(1, c, f"required column '{c}' is missing") for c in missing
            ]

        rows: list[RawRow] = []
        for sheet_row, values in enumerate(rows_iter, start=2):
            cells = {
                key: _to_str(values[idx] if idx < len(values) else None)
                for key, idx in col_index.items()
            }
            if all(v == "" for v in cells.values()):
                continue  # skip a fully-blank row
            rows.append(RawRow(row_number=sheet_row, cells=cells))
        return rows, []
    finally:
        wb.close()


# --------------------------------------------------------------- validation

def structural_row_errors(rows: list[RawRow]) -> list[RowError]:
    """Every no-database problem across `rows`, sorted by (row_number, column).

    Covers required-non-empty, numeric/date well-formedness, and per-group field
    consistency. Collects ALL errors (never stops at the first) so the operator
    gets one complete report.
    """
    errors: list[RowError] = []
    for row in rows:
        _required_errors(row, errors)
        _numeric_errors(row, errors)
        _date_errors(row, errors)
    _group_consistency_errors(rows, errors)
    errors.sort(key=lambda e: (e.row_number, e.column))
    return errors


def _required_errors(row: RawRow, errors: list[RowError]) -> None:
    for col in REQUIRED_COLUMNS:
        if not row.cells.get(col, "").strip():
            errors.append(RowError(row.row_number, col, f"{col} is required"))


# A generous sanity ceiling on a single money cell (paise). This bounds ONE cell,
# not a batch total: a single 50,000-line group can sum well past int8 (5e19 >
# 9.22e18), so any per-column total is the service's concern — here we only reject
# an absurd/typo value (a huge amount/rate) at validation with a clear message
# instead of a deferred insert overflow at generate.
_MAX_MONEY_PAISE = 10**15  # = Rs 10,000,000,000,000

# `quantity` maps to a Numeric(14,3) column: at most 11 integer digits and 3 decimal
# places. A larger/more-precise value passes the `> 0` check but overflows the column
# at generate — BURNING an already-reserved statutory number — so bound it here as a
# validation row error (paralleling the _MAX_MONEY_PAISE money ceiling).
_MAX_QTY = Decimal(10) ** 11  # exclusive ceiling: 11 integer digits
_QTY_MAX_DECIMALS = 3


def _qty_out_of_column(qty: Decimal) -> bool:
    """True if a positive quantity won't fit the Numeric(14,3) column."""
    if qty >= _MAX_QTY:
        return True
    exponent = qty.as_tuple().exponent
    return isinstance(exponent, int) and -exponent > _QTY_MAX_DECIMALS


def _numeric_errors(row: RawRow, errors: list[RowError]) -> None:
    """Validate quantity (> 0, fits Numeric(14,3)), amount (>= 0), gst_rate
    (0..100), and PIN codes.

    Emptiness is left to `_required_errors`, so these fire only on a present-but-
    malformed value (no duplicate error for a blank required cell). Pincodes are
    OPTIONAL, so they are checked only when present (must be exactly 6 digits).
    """
    c = row.cells
    if c.get("quantity", "").strip():
        qty = parse_qty(c["quantity"])
        if qty is None or qty <= 0:
            errors.append(RowError(row.row_number, "quantity",
                                   "quantity must be a positive number"))
        elif _qty_out_of_column(qty):
            errors.append(RowError(
                row.row_number, "quantity",
                "quantity is too large or too precise "
                "(max 11 digits before and 3 after the decimal point)"))
    if c.get("amount", "").strip():
        amount = parse_paise(c["amount"])
        if amount is None or amount < 0:
            errors.append(RowError(row.row_number, "amount",
                                   "amount must be a number >= 0"))
        elif amount > _MAX_MONEY_PAISE:
            errors.append(RowError(row.row_number, "amount", "amount is too large"))
    rate = c.get("rate", "").strip()
    if rate:  # rate may be free text; only a PARSED-numeric-and-out-of-range rate is wrong
        rate_paise = parse_paise(rate)
        if rate_paise is not None and rate_paise < 0:
            errors.append(RowError(row.row_number, "rate", "rate must not be negative"))
        elif rate_paise is not None and rate_paise > _MAX_MONEY_PAISE:
            errors.append(RowError(row.row_number, "rate", "rate is too large"))
    if c.get("gst_rate", "").strip():
        gst = parse_qty(c["gst_rate"])
        if gst is None or gst < 0 or gst > 100:
            errors.append(RowError(row.row_number, "gst_rate",
                                   "gst_rate must be a number between 0 and 100"))
    for col in ("ship_to_pincode", "consignee_pincode"):
        pin = c.get(col, "").replace(" ", "")
        # isascii() guards against non-ASCII digit code points (e.g. Devanagari
        # digits) that isdigit() alone would accept for a statutory PIN.
        if pin and not (len(pin) == 6 and pin.isascii() and pin.isdigit()):
            errors.append(RowError(row.row_number, col,
                                   f"{col} must be a 6-digit PIN code"))


def _date_errors(row: RawRow, errors: list[RowError]) -> None:
    text = row.cells.get("challan_date", "").strip()
    if text and parse_date(text) is None:
        errors.append(RowError(row.row_number, "challan_date",
                               "challan_date is not a recognisable date"))


def _group_consistency_errors(rows: list[RawRow], errors: list[RowError]) -> None:
    """Flag every row whose GROUP_CONSISTENT_FIELDS diverge from its group's first row.

    Rows are grouped by `challan_group`; a blank group is left to `_required_errors`.
    """
    groups: dict[str, list[RawRow]] = {}
    for row in rows:
        key = row.cells.get("challan_group", "").strip()
        if key:
            groups.setdefault(key, []).append(row)

    for key, members in groups.items():
        if len(members) < 2:
            continue
        first = members[0]
        for field_name in GROUP_CONSISTENT_FIELDS:
            baseline = first.cells.get(field_name, "")
            for row in members[1:]:
                if row.cells.get(field_name, "") != baseline:
                    errors.append(RowError(
                        row.row_number, field_name,
                        f"{field_name} must match the rest of group '{key}'"))
